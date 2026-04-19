"""PowerPoint slide generation using python-pptx."""

import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN


# Theme colors
COLOR_BG = RGBColor(0x1A, 0x1A, 0x2E)        # Dark navy
COLOR_ACCENT = RGBColor(0x16, 0x21, 0x3E)      # Slightly lighter navy
COLOR_HIGHLIGHT = RGBColor(0x0F, 0x3C, 0x78)   # Blue accent
COLOR_TITLE_TEXT = RGBColor(0xE0, 0xE0, 0xFF)  # Light lavender
COLOR_BODY_TEXT = RGBColor(0xCC, 0xCC, 0xCC)   # Light grey
COLOR_BULLET_ACCENT = RGBColor(0x4A, 0x9E, 0xFF)  # Bright blue for bullets


def _set_bg(slide, prs):
    """Fill slide background with dark navy."""
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = COLOR_BG


def _add_title_slide(prs: Presentation, title: str, subtitle: str) -> None:
    layout = prs.slide_layouts[6]  # Blank layout
    slide = prs.slides.add_slide(layout)
    _set_bg(slide, prs)

    w, h = prs.slide_width, prs.slide_height

    # Accent bar on left
    bar = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(0.12), h)
    bar.fill.solid()
    bar.fill.fore_color.rgb = COLOR_BULLET_ACCENT
    bar.line.fill.background()

    # Title text box
    tf_title = slide.shapes.add_textbox(Inches(0.4), Inches(2.2), w - Inches(0.8), Inches(1.6))
    frame = tf_title.text_frame
    frame.word_wrap = True
    p = frame.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = title
    run.font.size = Pt(44)
    run.font.bold = True
    run.font.color.rgb = COLOR_TITLE_TEXT

    # Subtitle
    tf_sub = slide.shapes.add_textbox(Inches(0.4), Inches(4.0), w - Inches(0.8), Inches(0.8))
    frame2 = tf_sub.text_frame
    p2 = frame2.paragraphs[0]
    run2 = p2.add_run()
    run2.text = subtitle or "AI-Generated Presentation"
    run2.font.size = Pt(20)
    run2.font.color.rgb = COLOR_BULLET_ACCENT
    run2.font.italic = True


def _add_content_slide(prs: Presentation, title: str, bullets: list[str], notes: str = "") -> None:
    layout = prs.slide_layouts[6]  # Blank layout
    slide = prs.slides.add_slide(layout)
    _set_bg(slide, prs)

    w, h = prs.slide_width, prs.slide_height

    # Header bar
    header = slide.shapes.add_shape(1, Inches(0), Inches(0), w, Inches(1.1))
    header.fill.solid()
    header.fill.fore_color.rgb = COLOR_HIGHLIGHT
    header.line.fill.background()

    # Slide title in header
    tf_title = slide.shapes.add_textbox(Inches(0.3), Inches(0.1), w - Inches(0.6), Inches(0.9))
    frame = tf_title.text_frame
    p = frame.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = title
    run.font.size = Pt(28)
    run.font.bold = True
    run.font.color.rgb = COLOR_TITLE_TEXT

    # Bullets content area
    tf_body = slide.shapes.add_textbox(Inches(0.5), Inches(1.3), w - Inches(0.8), h - Inches(1.6))
    body_frame = tf_body.text_frame
    body_frame.word_wrap = True

    for i, bullet in enumerate(bullets):
        p = body_frame.paragraphs[0] if i == 0 else body_frame.add_paragraph()
        p.space_before = Pt(6)
        p.space_after = Pt(2)
        run = p.add_run()
        run.text = f"•  {bullet}"
        run.font.size = Pt(18)
        run.font.color.rgb = COLOR_BODY_TEXT

    # Speaker notes
    if notes:
        slide.notes_slide.notes_text_frame.text = notes


def _add_summary_slide(prs: Presentation, summary: str) -> None:
    layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(layout)
    _set_bg(slide, prs)

    w, h = prs.slide_width, prs.slide_height

    # Header bar
    header = slide.shapes.add_shape(1, Inches(0), Inches(0), w, Inches(1.1))
    header.fill.solid()
    header.fill.fore_color.rgb = COLOR_HIGHLIGHT
    header.line.fill.background()

    tf_title = slide.shapes.add_textbox(Inches(0.3), Inches(0.1), w - Inches(0.6), Inches(0.9))
    frame = tf_title.text_frame
    p = frame.paragraphs[0]
    run = p.add_run()
    run.text = "Summary"
    run.font.size = Pt(28)
    run.font.bold = True
    run.font.color.rgb = COLOR_TITLE_TEXT

    tf_body = slide.shapes.add_textbox(Inches(0.5), Inches(1.4), w - Inches(1.0), h - Inches(1.8))
    body_frame = tf_body.text_frame
    body_frame.word_wrap = True
    p = body_frame.paragraphs[0]
    run = p.add_run()
    run.text = summary
    run.font.size = Pt(16)
    run.font.color.rgb = COLOR_BODY_TEXT


def generate_pptx(presentation_data: dict, output_path: str) -> str:
    """
    Generate a PowerPoint file from structured presentation data.
    Returns the output file path.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)

    _add_title_slide(prs, presentation_data["title"], presentation_data.get("subtitle", ""))

    for slide_data in presentation_data.get("slides", []):
        _add_content_slide(
            prs,
            slide_data["title"],
            slide_data.get("bullets", []),
            slide_data.get("notes", ""),
        )

    if presentation_data.get("summary"):
        _add_summary_slide(prs, presentation_data["summary"])

    prs.save(output_path)
    return output_path
